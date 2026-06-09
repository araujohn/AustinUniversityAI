"""
Create 'SuperKart Presentation - Final v2.pptx' from the v1 deck by embedding the
top-4 EDA charts (extracted from the notebook) into slide 5, with trimmed bullets.
"""
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from PIL import Image

SRC = "SuperKart Presentation - Final.pptx"          # v1 (already regenerated)
OUT = "SuperKart Presentation - Final v2.pptx"
DARK = RGBColor(0x33, 0x33, 0x33)

prs = Presentation(SRC)
s = prs.slides[4]  # Slide 5: EDA Results

# --- locate body placeholder (idx 1) and shrink to a left column ---
body = None
for ph in s.placeholders:
    if ph.placeholder_format.idx == 1:
        body = ph
        break
body.left, body.top, body.width, body.height = (
    Inches(0.22), Inches(1.0), Inches(3.05), Inches(3.9))

def set_body(tf, items, base_size=11, sub_size=10):
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    first = True
    for text, level in items:
        para = p0 if first else tf.add_paragraph()
        first = False
        para.level = level
        run = para.add_run()
        run.text = text
        run.font.size = Pt(base_size if level == 0 else sub_size)
        run.font.color.rgb = DARK
        para.space_after = Pt(6)

set_body(body.text_frame, [
    ("Target sale averages ₹3,464, roughly symmetric.", 0),
    ("Top drivers: Product_MRP (0.79) and Product_Weight (0.74); Allocated_Area ≈ 0.", 0),
    ("Store format: Departmental > Type1 > Type2 > Food Mart; Tier-1 and High-size lead.", 0),
    ("Revenue: OUT004 ≫ OUT003 > OUT001 > OUT002.", 0),
])

# --- 2x2 chart grid on the right ---
def place_img(path, box_l, box_t, box_w, box_h):
    iw, ih = Image.open(path).size
    ar = iw / ih
    if box_w / box_h > ar:      # box is wider than image -> fit by height
        h = box_h; w = h * ar
    else:                        # fit by width
        w = box_w; h = w / ar
    l = box_l + (box_w - w) / 2
    t = box_t + (box_h - h) / 2
    s.shapes.add_picture(path, Inches(l), Inches(t), Inches(w), Inches(h))

cell_w, cell_h = 3.15, 1.90
x0, x1 = 3.45, 6.72
y0, y1 = 1.02, 3.02
place_img("eda_imgs/corr_heatmap.png",    x0, y0, cell_w, cell_h)  # top-left
place_img("eda_imgs/target_dist.png",     x1, y0, cell_w, cell_h)  # top-right
place_img("eda_imgs/sales_storetype.png", x0, y1, cell_w, cell_h)  # bottom-left
place_img("eda_imgs/revenue_outlet.png",  x1, y1, cell_w, cell_h)  # bottom-right

prs.save(OUT)
print("Saved", OUT, "with", len(prs.slides._sldIdLst), "slides; slide 5 has",
      sum(1 for sh in s.shapes if sh.shape_type == 13), "pictures")
